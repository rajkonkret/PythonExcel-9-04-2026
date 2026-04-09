import openpyxl

wb = openpyxl.load_workbook('videogamesales.xlsx')
ws = wb.active

print(wb)  # <openpyxl.workbook.workbook.Workbook object at 0x000001D288501550>
print(ws)  # <Worksheet "vgsales">

ws = wb['vgsales']  # wybieramy arkusz do działąń

print("Total number of rows:", ws.max_row)  # Total number of rows: 16328
print("Total number of columns:", ws.max_column)  # Total number of columns: 10

print("Value in cell A1 is:", ws['A1'].value)  # Value in cell A1 is: Rank
