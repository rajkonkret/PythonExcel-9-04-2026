# Główne funkcje:
# Odczyt/Zapis: Odczytywanie danych z istniejących plików i tworzenie nowych.
# Modyfikacja: Edycja istniejących arkuszy, zmiana wartości komórek, dodawanie formuł.
# Formatowanie: Zmiana stylów komórek, obramowań, kolorów i czcionek.
# Wykresy: Tworzenie wykresów na podstawie danych w arkuszu.

# pip
from openpyxl import Workbook, load_workbook

# pip install openpyxl
# pip list

wb = Workbook()  # tworzy szablon pliku excel
ws = wb.active  # ustawiamy arkusz w pliku

ws['A1'] = 42

# zapisanie arkusza do pliku excel
wb.save("sample.xlsx")
wb.close()
